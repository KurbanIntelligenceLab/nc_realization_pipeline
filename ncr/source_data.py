"""Source Data workbooks: one Excel file per figure, one sheet per panel, as Nature Portfolio
journals ask. Every value is taken from figure_data.json, the same export the figures are drawn
from. Wherever a plotted quantity follows from counts in the sheet (shares, Wilson intervals,
coverage bounds, years from release, counts by outcome) it is written as an Excel formula, so a
reader can audit it; values the sheet cannot rederive (the label-effect intervals, median ages,
standardized differences) are stored with their source named in the README sheet."""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
Z95 = "NORMSINV(0.975)"


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows():
        for c in row:
            c.font = BOLD if c.row == 1 else FONT
    for j, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(48, len(str(h)) + 2))
    ws.freeze_panes = "A2"
    return ws


def _readme(wb, title, lines):
    ws = wb.active
    ws.title = "README"
    ws.append([title])
    ws["A1"].font = BOLD
    for ln in lines:
        ws.append([ln])
        ws.cell(row=ws.max_row, column=1).font = FONT
    ws.column_dimensions["A"].width = 120


def read_audit(path: str) -> list:
    """The literature audit of individual outcomes (Supplementary Note 6), if released; else empty."""
    import csv
    if not path or not os.path.exists(path):
        return []
    with open(path, newline="") as fh:
        return list(csv.DictReader(fh))


def _audit_sheet(wb, audit, keep):
    rows = [[a["composition"], a["release"], a["in_cohort_10yr"], a["threshold_side"], a["record_id"],
             a["record_source"], a["record_date"], a["finding"], a["what_the_source_establishes"],
             a["classification_under_outcome_definition"],
             a["classification_if_any_earlier_experimental_report_counts"], a["evidence"]]
            for a in audit if keep(a)]
    _sheet(wb, "Literature_audit", ["Composition", "Release", "Among ten-year cohort realizations",
                                    "Threshold side (+/-25 meV/atom, five years)", "Earliest matching record",
                                    "Record source", "Record date", "Literature finding",
                                    "What the source establishes", "Classification under the outcome definition",
                                    "Classification if any earlier experimental report counts", "Evidence"], rows)


def _wilson(k, n, z):
    """Wilson score bounds as Excel formulas for counts in cells k and n, with critical value z."""
    core = f"({k}/{n}+{z}^2/(2*{n}))/(1+{z}^2/{n})"
    half = f"{z}*SQRT({k}/{n}*(1-{k}/{n})/{n}+{z}^2/(4*{n}^2))/(1+{z}^2/{n})"
    return f"=MAX(0,{core}-{half})", f"=MIN(1,{core}+{half})"


def fig1(fd, path, audit=()):
    led = fd["fig1"]["ledger"]
    from .util import decimal_year, parse_date
    marks = [(a["release"], decimal_year(parse_date(a["record_date"])[0]), f'{a["composition"]}: {a["finding"]}')
             for a in audit if a.get("in_cohort_10yr") == "yes"]

    def mark(release, year):
        # offsets in the ledger are rounded to 1e-4 yr; records of one release are months apart
        hits = [m for rel, y, m in marks if rel == release and abs(y - year) < 1e-3]
        return hits[0] if len(hits) == 1 else ""
    order = sorted(led, key=lambda k: led[k]["date"])
    wb = Workbook()
    _readme(wb, "Source Data for Fig. 1", [
        "Sheet Fig1_records: every plotted point, one row per prediction already known at release or "
        "first recorded within its longest horizon of complete follow-up.",
        "  Earliest record is the decimal year of the earliest matching experimental structure record in the "
        "analyzed sources "
        "(year-only dates placed on 1 July); Years from release is a formula.",
        "Sheet Fig1_outcomes: the numbers printed in the figure. Counts are formulas over Fig1_records; "
        "median ages are computed by the released code (ncr/analysis.py, _ledger). Not recorded within "
        "follow-up includes predictions first recorded only after their complete follow-up, which no "
        "analysis counts.",
        "All dates are those of the earliest matching record in the analyzed sources (COD and the ICSD-derived "
        "records). Column Literature audit marks the four OQMD v1.0 compositions recorded within ten years, which "
        "were checked individually against the literature after the outcomes were measured; sheet "
        "Literature_audit gives what each source establishes. The counts remain those of the database matching.",
    ][: None if audit else -1])
    rows, r = [], 2
    for k in order:
        v = led[k]
        for y in sorted(v["known_years"]):
            rows.append([k, v["date"], "Already known", y, f"=D{r}-B{r}", ""])
            r += 1
        for off in sorted(v["later_offsets"]):
            rows.append([k, v["date"], "Recorded later", v["date"] + off, f"=D{r}-B{r}",
                         mark(k, v["date"] + off)])
            r += 1
    _sheet(wb, "Fig1_records", ["Release", "Release date (decimal year)", "Outcome",
                                "Earliest record (decimal year)", "Years from release",
                                "Literature audit (Supplementary Note 6)"], rows)
    out = []
    for i, k in enumerate(order, start=2):
        v = led[k]
        out.append([k, v["n"], f'=COUNTIFS(Fig1_records!A:A,A{i},Fig1_records!C:C,"Already known")',
                    f'=COUNTIFS(Fig1_records!A:A,A{i},Fig1_records!C:C,"Recorded later")', f"=B{i}-C{i}-D{i}",
                    v.get("age_median"), v["t_max"]])
    _sheet(wb, "Fig1_outcomes", ["Release", "Predictions", "Already known", "Recorded within follow-up", "Not recorded within follow-up",
                                 "Median age of already-known records (years)",
                                 "Longest horizon of complete follow-up (years)"], out)
    if audit:
        _audit_sheet(wb, audit, lambda a: a.get("in_cohort_10yr") == "yes")
    wb.save(path)


def fig2(fd, path, audit=()):
    wb = Workbook()
    _readme(wb, "Source Data for Fig. 2", [
        "Fig2a_coverage: capture-recapture coverage for records dated before 2017. Estimate = records in "
        "the source / Chapman total; the interval uses Chapman total +/- 1.96 standard errors, with the total "
        "never below the observed union. Chapman total, its standard error and the union are computed by the "
        "released code (ncr/analysis.py).",
        "Fig2b_already_known: counts by source and the share already known at release, with a Wilson 95% "
        "interval (formulas).",
        "Fig2c_recorded_later: share recorded within each complete horizon, with Wilson intervals simultaneous "
        "across the horizons of each release (Bonferroni; formulas).",
    ] + ([
        "Counts are database-matched: dates are those of the earliest matching record in the analyzed sources. "
        "A literature audit found earlier experimental reports for two of the four OQMD v1.0 compositions "
        "recorded within ten years (Supplementary Note 6; marked in Source Data Fig. 1).",
    ] if audit else []))
    rows, r = [], 2
    for c in (fd.get("fig2") or {}).get("coverage", []):
        for part, lab in (("COD", "COD"), ("ICSD-derived", "ICSD-derived"), ("either", "Either")):
            rows.append([c["source"], lab, c[part]["k"], c["chapman_total"], c["chapman_se"], c["union"],
                         f"=C{r}/D{r}", f"=C{r}/(D{r}+1.96*E{r})", f"=MIN(1,C{r}/MAX(D{r}-1.96*E{r},F{r}))"])
            r += 1
    _sheet(wb, "Fig2a_coverage", ["Release", "Source", "Records in source", "Chapman total", "Chapman SE",
                                  "Union of sources", "Coverage", "Lower bound", "Upper bound"], rows)
    parts = {q["source"]: q for q in fd["fig2c"]}
    rows = []
    for i, a in enumerate(fd["fig2a"], start=2):
        d = a["already_known"]
        p = parts[a["source"]]
        lo, hi = _wilson(f"C{i}", f"B{i}", Z95) if d["n"] else (None, None)
        rows.append([a["source"], d["n"], d["k"], p["COD only"], p["both"], p["ICSD-derived only"],
                     f"=C{i}/B{i}" if d["n"] else None, lo, hi])
    _sheet(wb, "Fig2b_already_known", ["Release", "Predictions", "Already known", "COD only", "Both",
                                       "ICSD-derived only", "Share already known", "Wilson 95% lower",
                                       "Wilson 95% upper"], rows)
    rz = (fd.get("fig2") or {}).get("realized_by_horizon", [])
    m = {}
    for x in rz:
        m[x["source"]] = m.get(x["source"], 0) + 1
    rows = []
    for i, x in enumerate(sorted(rz, key=lambda q: (q["source"], q["t"])), start=2):
        z = f"NORMSINV(1-0.05/(2*E{i}))"
        lo, hi = _wilson(f"C{i}", f"D{i}", z) if x["n"] else (None, None)
        rows.append([x["source"], x["t"], x["k"], x["n"], m[x["source"]], f"=C{i}/D{i}" if x["n"] else None, lo, hi])
    _sheet(wb, "Fig2c_recorded_later", ["Release", "Horizon (years)", "Recorded later", "Predictions",
                                        "Horizons of this release (Bonferroni)", "Share recorded later",
                                        "Simultaneous 95% lower", "Simultaneous 95% upper"], rows)
    wb.save(path)


def fig3(fd, path, audit=()):
    d = fd["fig3"]
    wb = Workbook()
    _readme(wb, "Source Data for Fig. 3", [
        "Panel a is a schematic and has no source data.",
        "Fig3b_premise: compositions in both OQMD and the Materials Project and the number whose stable labels "
        "disagree, by distance from the OQMD threshold; share and Wilson 95% interval are formulas.",
        "Fig3c_threshold: entries and structures recorded within five years in each 5 meV/atom bin of the "
        "decomposition energy relative to the threshold.",
        "Fig3d_label_effect: blocked differences and their normal-approximation 95% intervals, computed by the "
        "released code (ncr/analysis.py); unreliable with this few events. The last row is the smallest effect "
        "detectable with 80% power in an approximate normal-theory calculation; it describes the design, not the "
        "effects the data exclude.",
    ] + ([
        "Events are database-matched. Sheet Literature_audit gives the literature check of the four threshold "
        "events (three labeled stable, one unstable) and what each source establishes; a sensitivity analysis "
        "that counts any earlier experimental report is computed by scripts/literature_sensitivity.py.",
    ] if audit else []))
    rows = []
    i = 2
    for p in d["premise"]:
        for b in p["bins"]:
            lo_b, hi_b = b["bin"]
            hi_b = "above" if hi_b is None or hi_b >= 1e8 else hi_b
            if not b.get("n") or b.get("rate") is None:
                rows.append([p["pair"], lo_b, hi_b, b.get("n") or 0, None, None, None, None])
            else:
                lo, hi = _wilson(f"E{i}", f"D{i}", Z95)
                rows.append([p["pair"], lo_b, hi_b, b["n"], round(b["rate"] * b["n"]), f"=E{i}/D{i}", lo, hi])
            i += 1
    _sheet(wb, "Fig3b_premise", ["Pair", "Distance from threshold, from (meV/atom)", "to (meV/atom)",
                                 "Compositions", "Labels disagree", "Share disagreeing", "Wilson 95% lower",
                                 "Wilson 95% upper"], rows)
    rows = [[b["center"], "Stable" if b["center"] < 0 else "Unstable", b["n"], b["k"]] for b in d["bins"]]
    _sheet(wb, "Fig3c_threshold", ["Bin centre (meV/atom, relative to threshold)", "Side", "Entries",
                                   "Structures recorded within five years"], rows)
    lab = {"primary": "primary", "window": "window", "placebo": "placebo"}

    def pp(v):
        return None if v is None else 100 * v
    rows = [[r_["label"], lab.get(r_["group"], r_["group"]), pp(r_.get("est")), pp(r_.get("lo")), pp(r_.get("hi"))]
            for r_ in d["forest"] if r_["group"] in lab]
    if d.get("mde") is not None:
        rows.append(["Smallest effect detectable with 80% power (approximate, normal theory), primary analysis", "reference",
                     pp(d["mde"]), None, None])
    _sheet(wb, "Fig3d_label_effect", ["Analysis", "Group", "Effect (percentage points)", "95% lower",
                                      "95% upper"], rows)
    if audit:
        _audit_sheet(wb, audit, lambda a: a.get("threshold_side") in ("stable", "unstable"))
    wb.save(path)


def figS1(fd, path):
    d = fd["fig1"]
    wb = Workbook()
    _readme(wb, "Source Data for Supplementary Fig. 1", [
        "Releases: release date and longest horizon of complete follow-up; the end of follow-up is a formula.",
        "Sources: last publication year covered by each source; all continue before 2010.",
        "Cutoff: the data cutoff and the start of the 1-year reporting-lag buffer (formula).",
    ])
    rows = [[r["label"], r["n"], r["date"], r["t_max"], f"=C{i}+D{i}"] for i, r in enumerate(d["releases"], 2)]
    _sheet(wb, "Releases", ["Release", "Predictions", "Release date (decimal year)",
                            "Longest complete horizon (years)", "End of complete follow-up"], rows)
    _sheet(wb, "Sources", ["Source", "Last publication year covered"], [[s["name"], s["last_year"]]
                                                                         for s in d["sources"]])
    _sheet(wb, "Cutoff", ["Data cutoff (decimal year)", "Start of reporting-lag buffer"], [[d["cutoff"], "=A2-1"]])
    wb.save(path)


def figS2(fd, path):
    bal = fd["fig3"]["balance"]
    wb = Workbook()
    _readme(wb, "Source Data for Supplementary Fig. 2", [
        "Standardized differences (stable minus unstable side) by window, covariate and blocking scheme, "
        "computed by the released code (ncr/checks.py) before any outcome was measured. The pre-specified "
        "limit is 0.1 in absolute value.",
    ])
    rows = []
    for w in sorted(bal.get("windows", {}), key=float):
        for ser in bal["windows"][w]:
            for cv, v in ser["smd"].items():
                rows.append([float(w), cv, ser["label"], v])
    _sheet(wb, "SuppFig2_balance", ["Window (+/- meV/atom)", "Covariate", "Blocking scheme",
                                    "Standardized difference"], rows)
    wb.save(path)


def write_all(fd: dict, outdir: str, audit=()) -> list:
    os.makedirs(outdir, exist_ok=True)
    out = []
    for name, fn in (("Source_Data_Fig1.xlsx", fig1), ("Source_Data_Fig2.xlsx", fig2),
                     ("Source_Data_Fig3.xlsx", fig3), ("Source_Data_Supplementary_Fig1.xlsx", figS1),
                     ("Source_Data_Supplementary_Fig2.xlsx", figS2)):
        p = os.path.join(outdir, name)
        fn(fd, p, audit) if fn in (fig1, fig2, fig3) else fn(fd, p)
        out.append(p)
    return out
