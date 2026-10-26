"""End-to-end and unit tests. The end-to-end tests build a synthetic world with known truth
(tests/synthetic.py), run every pipeline step through the command line, and check the
estimates against quantities computed directly from the synthetic inputs."""
from __future__ import annotations

import gzip
import json
import math
import os
import subprocess
import sys

import numpy as np
import pytest

from ncr import stats as S
from ncr.snapshots import compute_dhd
from ncr.util import decimal_year, parse_date, read_csv, read_json
from synthetic import make

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STEPS = ("check", "hull", "reconstruction", "cohorts", "power", "premise", "coverage-known", "match",
         "estimate", "robustness", "figures")


def ncr(*args, config=None):
    cmd = [sys.executable, "-m", "ncr"] + (["--config", config] if config else []) + list(args)
    p = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True)
    assert p.returncode == 0, f"{' '.join(args)} failed:\n{p.stderr[-2000:]}"
    return p


def _y(s):
    d, _ = parse_date(s)
    return decimal_year(d) if d else math.nan


@pytest.fixture(scope="module")
def world(tmp_path_factory):
    d = str(tmp_path_factory.mktemp("world"))
    make(d)
    ncr("merge-experiments", "--in", os.path.join(d, "cod.csv"), os.path.join(d, "mp_icsd.csv"),
        "--out", os.path.join(d, "experiments.csv"))
    cfg = os.path.join(d, "config.json")
    for step in STEPS:
        ncr(step, config=cfg)
    return d


def _w(world, *p):
    return os.path.join(world, "work", *p)


# --------------------------------------------------------------------------- end-to-end
def test_every_step_writes_its_outputs(world):
    for p in (("checks", "reconstruction.json"), ("checks", "power.json"), ("checks", "premise.json"),
              ("checks", "coverage_known.json"), ("results", "summary.json"), ("results", "robustness.json"),
              ("figures", "fig1.pdf"), ("figures", "fig2.pdf"), ("figures", "fig3.pdf"), ("figures", "figS1.pdf"),
              ("figures", "figS2.pdf"), ("results", "figure_data.json"),
              ("manifest.json",)):
        assert os.path.exists(_w(world, *p)), p


def test_manifest_records_relative_paths_and_checksums(world):
    ev = [e for e in read_json(_w(world, "manifest.json"))["events"] if e["kind"] == "inputs"][-1]
    for rec in ev["files"].values():
        assert not os.path.isabs(rec["path"]) and len(rec["sha256"]) == 64


def test_reconstruction_agrees_with_displayed_labels(world):
    rec = read_json(_w(world, "checks", "reconstruction.json"))
    for r in rec.values():
        assert r["recomputed_recall_of_displayed_stable"] == 1.0
        assert r["new_composition_overlap"]["jaccard"] == 1.0


def test_pre_existence_equals_direct_count(world):
    rows = read_csv(_w(world, "cohorts", "srcb.csv"))
    m = {x["pred_id"]: x for x in read_csv(_w(world, "matches", "srcb.csv"))}
    direct = sum(1 for r in rows if m[r["pred_id"]]["db_first_T3"] and _y(m[r["pred_id"]]["db_first_T3"]) <= _y(r["entry_date"]))
    pe = read_json(_w(world, "results", "summary.json"))["base"]["srcb"]["pre_existence_all"]
    assert pe["k"] == direct and pe["n"] == len(rows)


def test_matches_keep_each_source_separately(world):
    m = read_csv(_w(world, "matches", "srcb.csv"))
    assert {"db_first_T3__COD", "db_first_T3__MP-ICSD"} <= set(m[0])
    for x in m:                                    # the overall first record is the earliest of the sources
        ds = [_y(x[k]) for k in ("db_first_T3__COD", "db_first_T3__MP-ICSD") if x[k]]
        assert (not ds and not x["db_first_T3"]) or abs(_y(x["db_first_T3"]) - min(ds)) < 1e-9


def test_theoretical_records_are_never_matched(world):
    for x in read_csv(_w(world, "matches", "srcb.csv")):
        assert x["db_id_T3"] != "ctheo"


def test_threshold_effect_equals_direct_blocked_difference(world):
    rows = read_csv(_w(world, "cohorts", "threshold__srcb.csv"))
    m = {x["pred_id"]: x for x in read_csv(_w(world, "matches", "threshold__srcb.csv"))}
    y, xs, side, blk = [], [], [], []
    for r in rows:
        e, f = _y(r["entry_date"]), _y(m[r["pred_id"]]["db_first_T3"])
        y.append(float(np.isfinite(f) and e < f <= e + 5))
        xs.append(float(r["x_mev"])); side.append(r["side"]); blk.append(r["block"])
    y, xs, side, blk = map(np.array, (y, xs, side, blk))
    inw = np.abs(xs) <= 25
    num, tot = 0.0, 0
    for b in np.unique(blk[inw]):
        A, B = inw & (blk == b) & (side == "A"), inw & (blk == b) & (side == "B")
        if A.sum() >= 2 and B.sum() >= 2:
            num += (A.sum() + B.sum()) * (y[A].mean() - y[B].mean())
            tot += A.sum() + B.sum()
    eff = read_json(_w(world, "results", "summary.json"))["threshold"]["threshold__srcb"]["windows"]["25"]["effect"]
    assert abs(eff["estimate"] - num / tot) < 1e-12 and eff["se"] > 0


def test_label_effect_is_recovered(world):
    eff = read_json(_w(world, "results", "summary.json"))["threshold"]["threshold__srcb"]["windows"]["50"]["effect"]
    assert eff["estimate"] > 0.05 and eff["ci_low"] > 0


def test_capture_recapture_is_consistent(world):
    cr = read_json(_w(world, "results", "summary.json"))["capture_recapture"]["srcb"]
    assert cr["chapman_total"] >= cr["union"] and 0 < cr["coverage_union"] <= 1


def test_power_reports_every_blocking_scheme(world):
    pw = read_json(_w(world, "checks", "power.json"))["SrcB"]
    assert set(pw["block_scheme_balance_by_window"]) == {"10", "25", "50"}
    assert set(pw["block_scheme_candidates_25meV"]["schemes"]) == {"class", "class_en", "class_en_nel"}


def test_premise_check_runs(world):
    pr = read_json(_w(world, "checks", "premise.json"))["srcb vs srca_r1"]
    assert pr["matched"] > 0 and len(pr["bins"]) >= 5


def test_known_compound_coverage_uses_only_experimental_compounds(world):
    cov = read_json(_w(world, "checks", "coverage_known.json"))
    assert cov["srcb"]["icsd_origin_compositions"] > 0


def test_robustness_variants_are_nested(world):
    rob = read_json(_w(world, "results", "robustness.json"))
    get = lambda v: v["srcb"]["srcb::srcb"]["pre_existence"]["k"]  # noqa: E731
    assert get(rob["rematch"]["composition_tolerance_halved"]) <= get(rob["primary"]) \
        <= get(rob["rematch"]["composition_tolerance_doubled"])
    assert "local_linear_25meV" in rob["reanalysis"]["threshold__srcb"]


def test_figures_pass_their_presentation_checks(world):
    checks = read_json(_w(world, "results", "figure_checks.json"))
    assert checks and all(checks.values())
    fd = read_json(_w(world, "results", "figure_data.json"))["fig3"]
    assert fd["premise"] and fd["balance"]["series"] and fd["forest"][0]["group"] == "primary"


# --------------------------------------------------------------------------- converters
def test_convert_oqmd(tmp_path):
    H = ["entry_id", "composition_id", "path", "label", "duplicate_of_id", "prototype_id", "ntypes", "natoms",
         "entry_delta_e", "entry_stability", "fe_id", "fit_id", "fe_delta_e", "fe_stability"]
    P = "/home/oqmd/libraries/"
    rows = [
        [1, "Li1", P + "prototypes/elements/A2/Li", "A2", "NULL", "A2", 1, 2, "NULL", "NULL", 11, "standard", 0.004, 0.004],
        [2, "O1", P + "prototypes/elements/A1/O", "A1", "NULL", "A1", 1, 2, "NULL", "NULL", 12, "standard", 0.0, 0.0],
        [3, "Fe1", P + "prototypes/elements/A2/Fe", "A2", "NULL", "A2", 1, 2, "NULL", "NULL", 13, "standard", 0.0, 0.0],
        [4, "Li2 O1", P + "icsd/101", "icsd-101", 4, "C1", 2, 3, "NULL", "NULL", 14, "standard", -2.0, -0.05],
        [5, "Li2 O1", P + "prototypes/binaries/B1/Li_O", "NULL", "NULL", "B1", 2, 3, "NULL", "NULL", 15, "standard", -1.9, 0.10],
        [6, "Fe1 Li1 O1", P + "prototypes/ternaries/half_heusler/Fe/Li_O", "NULL", "NULL", "NULL", 3, 3, "NULL", "NULL", 16, "standard", -1.0, 0.02],
        [6, "Fe1 Li1 O1", P + "prototypes/ternaries/half_heusler/Fe/Li_O", "NULL", "NULL", "NULL", 3, 3, "NULL", "NULL", 26, "standard", -1.6, -0.03],
        [7, "Fe1 Li1 O1", P + "prototypes/ternaries/heusler/Fe/Li_O", "NULL", 6, "NULL", 3, 4, "NULL", "NULL", 17, "standard", -1.6, -0.03],
        [8, "Fe1 O1", P + "prototypes/binaries/B1/Fe_O", "NULL", "NULL", "B1", 2, 2, "NULL", "NULL", 18, "standard", -1.3, "NULL"],
        [9, "Fe2 O3", P + "prototypes/binaries/D5_1/Fe_O", "NULL", "NULL", "D5_1", 2, 5, "NULL", "NULL", "NULL", "NULL", "NULL", "NULL"],
        [10, "Fe1 Li2", "/home/oqmd/test/tmp/x", "NULL", "NULL", "NULL", 2, 3, "NULL", "NULL", 20, "standard", 0.2, 0.2],
        [11, "Fe3 O4", P + "icsd/202", "icsd-202", 11, "H1_1", 2, 14, "NULL", "NULL", "NULL", "NULL", "NULL", "NULL"],
        [12, "Fe3 O4", P + "prototypes/binaries/H1_1/Fe_O", "NULL", "NULL", "H1_1", 2, 7, "NULL", "NULL", 22, "standard", -1.0, 0.5],
        [13, "Li1 O2", P + "prototypes/binaries/C2/Li_O", "NULL", "NULL", "C2", 2, 3, "NULL", "NULL", 23, "standard", -1.5, -0.01],
        [13, "Li1 O2", P + "prototypes/binaries/C2/Li_O", "NULL", "NULL", "C2", 2, 3, "NULL", "NULL", 33, "standard", -1.2, 0.29],
    ]
    e, s, out = tmp_path / "e.tsv", tmp_path / "s.tsv", tmp_path / "snap.csv"
    e.write_text("\t".join(H) + "\n" + "".join("\t".join(str(x) for x in r) + "\n" for r in rows))
    s.write_text("entry_id\tstructure_id\tlabel\tspacegroup_id\n6\t1\tinput\t216\n6\t2\tstandard\t225\n")
    ncr("convert-oqmd", "--entries", str(e), "--structures", str(s), "--out", str(out))
    rep = read_json(str(out).replace(".csv", "_conversion_report.json"))
    snap = {r["entry_id"]: r for r in read_csv(str(out))}
    assert rep["dropped_duplicates"] == 1 and rep["dropped_no_formation_energy"] == 2
    assert snap["oqmd-6"]["energy_per_atom"] == "-1.6"                 # lowest formation-energy row kept
    assert snap["oqmd-13"]["energy_per_atom"] == "-1.5" and snap["oqmd-13"]["displayed_stable"] == "1"
    assert snap["oqmd-12"]["source_known_composition"] == "1"          # ICSD entry without an energy still counts
    assert snap["oqmd-5"]["source_known_composition"] == "1" and snap["oqmd-4"]["is_hypothetical"] == "0"
    assert snap["oqmd-10"]["is_hypothetical"] == "0" and snap["oqmd-8"]["displayed_stable"] == ""
    assert snap["oqmd-6"]["spacegroup_number"] == "225" and "heusler" in snap["oqmd-6"]["prototype"]
    assert rep["displayed_stable_new_compositions"] == 2 and rep["elemental_references_added"] == 3


def test_convert_mp(tmp_path):
    B = tmp_path / "build"
    def rec(mid, comp, fe, stable, dhd, tt="GGA_GGA+U"):
        return {"material_id": mid, "thermo_type": tt, "composition_reduced": comp, "formation_energy_per_atom": fe,
                "is_stable": stable, "decomposition_enthalpy": dhd, "energy_above_hull": max(dhd, 0.0),
                "deprecated": False, "symmetry": {"number": 225}}
    thermo = [rec("mp-1", {"Li": 2.0, "O": 1.0}, -2.1, True, -0.30), rec("mp-2", {"Li": 1.0, "Fe": 1.0, "O": 1.0}, -1.55, False, 0.02),
              rec("mp-3", {"Li": 1.0, "O": 2.0}, -1.4, False, 0.05),
              rec("mp-1", {"Li": 2.0, "O": 1.0}, -2.0, True, -0.2, tt="GGA_GGA+U_R2SCAN"), rec("mp-4", {"Fe": 1.0}, 0.0, True, 0.0)]
    prov = [{"material_id": "mp-1", "theoretical": False, "created_at": {"$date": "2022-10-20T00:00:00Z"}, "database_IDs": {"icsd": ["icsd-1"]}},
            {"material_id": "mp-2", "theoretical": True, "created_at": {"$date": "2022-10-20T00:00:00Z"}, "database_IDs": {}},
            {"material_id": "mp-3", "theoretical": True, "created_at": {"$date": "2019-01-01T00:00:00Z"}, "database_IDs": {}}]
    (B / "thermo" / "x").mkdir(parents=True); (B / "provenance" / "x").mkdir(parents=True)
    with gzip.open(B / "thermo" / "x" / "a.jsonl.gz", "wt") as fh:
        fh.write("\n".join(json.dumps(r) for r in thermo) + "\n")
    with gzip.open(B / "provenance" / "x" / "p.jsonl.gz", "wt") as fh:
        fh.write("\n".join(json.dumps(r) for r in prov) + "\n")
    out = tmp_path / "mp.csv"
    ncr("convert-mp", "--build-dir", str(B), "--out", str(out))
    rep = read_json(str(out).replace(".csv", "_conversion_report.json"))
    snap = {r["entry_id"]: r for r in read_csv(str(out))}
    assert rep["chosen_materials"] == 4 and snap["mp-1"]["energy_per_atom"] == "-2.1"   # thermo type from the field
    assert snap["mp-1"]["origin_icsd_id"] == "icsd" and snap["mp-2"]["is_hypothetical"] == "1"
    assert snap["mp-2"]["displayed_dhd"] == "0.02" and snap["mp-4"]["is_hypothetical"] == "0"


def test_convert_cod(tmp_path):
    from ncr.cod import FIELDS
    rows = [{"file": "1", "formula": "- Fe3 O4 -", "year": "1986", "flags": "has coordinates"},
            {"file": "2", "formula": "- Li Fe O -", "year": "2019", "method": "single crystal",
             "flags": "has coordinates,has disorder", "cellpressure": "2000000"},
            {"file": "3", "formula": "- Li O2 -", "year": "2020", "method": "theoretical"},
            {"file": "4", "formula": "- Fe O -", "year": "2001", "status": "retracted"},
            {"file": "5", "formula": "", "calcformula": "- Li2 O -", "year": ""},
            {"file": "6", "formula": "- C6 H6 -", "year": "2005"}]
    p = tmp_path / "cod.tsv"
    p.write_text("\t".join(FIELDS) + "\n" + "".join("\t".join(r.get(k, "") for k in FIELDS) + "\n" for r in rows))
    out = tmp_path / "cod.csv"
    ncr("convert-cod", "--in", str(p), "--out", str(out))
    rep = read_json(str(out).replace(".csv", "_conversion_report.json"))
    ex = {r["exp_id"]: r for r in read_csv(str(out))}
    assert set(ex) == {"cod-1", "cod-2", "cod-6"}
    assert rep["excluded_theoretical"] == 1 and rep["excluded_retracted"] == 1 and rep["excluded_no_year"] == 1
    assert ex["cod-2"]["high_pressure"] == "1" and ex["cod-2"]["disordered"] == "1"
    from ncr.matching import read_experiments
    exps = read_experiments(str(out))
    assert str(exps[0]["date"]) == "1986-07-01"                       # year-only dates set to 1 July


def test_convert_icsd_derived_sources(tmp_path):
    t = tmp_path / "oqi.tsv"
    t.write_text("entry_id\tcomposition_id\tyear\n1\tFe3 O4\t2010\n2\tLi2 O1\tNULL\n3\tNa1 Cl1\t1995\n")
    ncr("convert-oqmd-icsd", "--in", str(t), "--out", str(tmp_path / "oqi.csv"))
    r1 = read_json(str(tmp_path / "oqi_conversion_report.json"))
    assert r1["rows_written"] == 2 and r1["year_max"] == 2010
    B = tmp_path / "b"
    (B / "provenance" / "x").mkdir(parents=True); (B / "thermo" / "x").mkdir(parents=True)
    prov = [{"material_id": "mp-1", "theoretical": False, "database_IDs": {"icsd": ["icsd-1"]},
             "references": ["@article{a, year = {2005}}", "@article{b, year = {1999}}"]},
            {"material_id": "mp-2", "theoretical": False, "database_IDs": {}, "references": []},
            {"material_id": "mp-3", "theoretical": True, "database_IDs": {}, "references": []}]
    therm = [{"material_id": "mp-1", "composition_reduced": {"Fe": 3.0, "O": 4.0}},
             {"material_id": "mp-2", "composition_reduced": {"Li": 2.0, "O": 1.0}}]
    with gzip.open(B / "provenance" / "x" / "p.jsonl.gz", "wt") as fh:
        fh.write("\n".join(json.dumps(x) for x in prov) + "\n")
    with gzip.open(B / "thermo" / "x" / "t.jsonl.gz", "wt") as fh:
        fh.write("\n".join(json.dumps(x) for x in therm) + "\n")
    ncr("convert-mp-icsd", "--build-dir", str(B), "--out", str(tmp_path / "mpi.csv"))
    r2 = read_json(str(tmp_path / "mpi_conversion_report.json"))
    mpi = read_csv(str(tmp_path / "mpi.csv"))
    assert r2["rows_written"] == 1 and r2["excluded_no_icsd_identifier"] == 1 and mpi[0]["first_report_date"] == "1999"


# --------------------------------------------------------------------------- units
def test_higher_energy_polymorph_is_above_the_hull():
    rows = [{"entry_id": "li", "formula": "Li", "energy_per_atom": 0.0, "is_hypothetical": False},
            {"entry_id": "o", "formula": "O", "energy_per_atom": 0.0, "is_hypothetical": False},
            {"entry_id": "gs", "formula": "Li2O", "energy_per_atom": -2.0, "is_hypothetical": True},
            {"entry_id": "poly", "formula": "Li2O", "energy_per_atom": -1.9, "is_hypothetical": True}]
    r = compute_dhd(rows, lambda x: x["is_hypothetical"])
    assert abs(r["poly"][0] - 0.1) < 1e-9 and r["gs"][0] < 0


def test_exact_identities():
    assert all(S.exact_checks().values())


def test_wilson_interval_is_exact_at_the_boundaries():
    assert S.wilson_ci(0, 136)[0] == 0.0 and S.wilson_ci(136, 136)[1] == 1.0


def test_binomial_test_is_exact_and_stable():
    assert abs(S.binomial_two_sided(0, 10) - 2 / 1024) < 1e-12
    assert abs(S.binomial_two_sided(2000, 4000) - 1.0) < 1e-9 and 0 < S.binomial_two_sided(1800, 4000) < 1e-8


def test_format_share_rounds_the_interval_outward():
    assert S.format_share(103 / 2860, *S.wilson_ci(103, 2860)) == "3.6% [2.9, 4.4]"


def test_minimum_detectable_effect_multiplier():
    assert abs(S.minimum_detectable_effect(0.05, 10**8, 10**8) / math.sqrt(2 * 0.05 * 0.95 / 10**8) - 2.8016) < 1e-3


def test_cutoff_in_the_future_is_rejected(tmp_path):
    from ncr import config as C
    p = tmp_path / "c.json"
    p.write_text(json.dumps({"cutoff_date": "2999-01-01", "releases": []}))
    with pytest.raises(C.ConfigError):
        C.cutoff(C.load(str(p)))


def test_blocked_design_simulation():
    sim = S.simulate_blocked_design(reps=400)
    assert abs(sim["mean_blocked"] - sim["true_effect"]) < 0.004
    assert sim["mean_unblocked"] > sim["true_effect"] + 0.03          # blocking removes confounding
    assert 0.93 <= sim["coverage"] <= 0.975


def test_source_data_workbooks_hold_every_plotted_record(world):
    from openpyxl import load_workbook
    sd = _w(world, "figures", "source_data")
    names = sorted(os.listdir(sd))
    assert names == sorted(["Source_Data_Fig1.xlsx", "Source_Data_Fig2.xlsx", "Source_Data_Fig3.xlsx",
                            "Source_Data_Supplementary_Fig1.xlsx", "Source_Data_Supplementary_Fig2.xlsx"])
    fd = read_json(_w(world, "results", "figure_data.json"))
    ws = load_workbook(os.path.join(sd, "Source_Data_Fig1.xlsx"))["Fig1_records"]
    want = sum(len(v["known_years"]) + len(v["later_offsets"]) for v in fd["fig1"]["ledger"].values())
    assert ws.max_row - 1 == want
    wb = load_workbook(os.path.join(sd, "Source_Data_Fig3.xlsx"))
    assert wb.sheetnames == ["README", "Fig3b_premise", "Fig3c_threshold", "Fig3d_label_effect"]
