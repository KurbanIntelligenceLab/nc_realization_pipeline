"""The literature audit marks exactly the audited Source Data rows and nothing without an audit file."""
import os

from openpyxl import load_workbook

from ncr import source_data


def _fd():
    # one release, two later records 1.211 and 2.2123 yr after 2014.2849315068493 (offsets rounded to 1e-4)
    return {"fig1": {"ledger": {"R (2014)": {"date": 2014.2849315068493, "n": 10, "known_years": [1990.5],
                                            "later_offsets": [1.211, 2.2123], "t_max": 10, "age_median": 24}}}}


def test_audit_marks_only_matching_rows(tmp_path):
    audit = [{"release": "R (2014)", "record_date": "2015-07-01", "composition": "X", "finding": "unverified",
              "in_cohort_10yr": "yes", "threshold_side": "stable", "record_id": "r1", "record_source": "COD",
              "what_the_source_establishes": "-", "classification_under_outcome_definition": "unchanged (event)",
              "classification_if_any_earlier_experimental_report_counts": "unchanged (event)", "evidence": "-"}]
    p = os.path.join(tmp_path, "f1.xlsx")
    source_data.fig1(_fd(), p, audit)
    wb = load_workbook(p)
    rows = list(wb["Fig1_records"].iter_rows(min_row=2, values_only=True))
    assert [r[5] or "" for r in rows] == ["", "X: unverified", ""]  # empty cells read back as None
    assert "Literature_audit" in wb.sheetnames


def test_no_audit_file_means_no_marks(tmp_path):
    assert source_data.read_audit(os.path.join(tmp_path, "missing.csv")) == []
    p = os.path.join(tmp_path, "f1.xlsx")
    source_data.fig1(_fd(), p)
    wb = load_workbook(p)
    assert "Literature_audit" not in wb.sheetnames
    assert all(not r[5] for r in wb["Fig1_records"].iter_rows(min_row=2, values_only=True))
